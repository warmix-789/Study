import openpyxl
from openpyxl.styles import Font, PatternFill

def create_test_excel(filename="company_schedule.xlsx"):
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "График работы"

    # Создаем шапку таблицы
    headers = ["ID", "ФИО", "Роль"] + [f"{i} ноя" for i in range(1, 8)]
    sheet.append(headers)

    # Делаем шапку жирной и заливаем серым цветом
    header_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill

    # Наши тестовые пользователи со статусами (Р - рабочий, В - выходной, О - отпуск, Б - больничный)
    test_data = [
        [1, "Иванов А.А.", "Сотрудник", "Р", "Р", "Р", "Р", "Р", "В", "В"],
        [2, "Петров Б.Б.", "Сотрудник", "Р", "Р", "В", "В", "Р", "Р", "Р"],
        [3, "Сидоров В.В.", "Менеджер",  "О", "О", "О", "О", "О", "В", "В"],
        [4, "Смирнова Г.Г.", "Сотрудник", "Б", "Б", "Б", "Р", "Р", "В", "В"],
        [5, "Черняев М.", "Админ",        "Р", "Р", "Р", "Р", "Р", "В", "В"]
    ]

    # Записываем данные в таблицу
    for row in test_data:
        sheet.append(row)

    # Расширяем колонку с ФИО
    sheet.column_dimensions['B'].width = 15

    wb.save(filename)
    print(f"Готово! Тестовый файл '{filename}' успешно создан.")

if __name__ == "__main__":
    create_test_excel()