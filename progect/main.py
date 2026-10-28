from fastapi import FastAPI, Depends, Request, Form, Response, Cookie, File, UploadFile
from fastapi.responses import HTMLResponse, RedirectResponse, StreamingResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from sqlalchemy.orm import Session
from typing import Optional
import hashlib
import models
import io
import openpyxl
from database import engine, SessionLocal

models.Base.metadata.create_all(bind=engine)
app = FastAPI(title="График Работы API")
templates = Jinja2Templates(directory="templates")

DAYS = [f"{i} ноя" for i in range(1, 31)]
TODAY_INDEX = 2 

def get_db():
    db = SessionLocal()
    try: yield db
    finally: db.close()

def hash_password(password: str): return hashlib.sha256(password.encode()).hexdigest()

@app.get("/manifest.json")
def get_manifest():
    return JSONResponse({
        "name": "HRM Система",
        "short_name": "График",
        "start_url": "/",
        "display": "standalone",
        "background_color": "#1c1c1e",
        "theme_color": "#1c1c1e",
        "icons": [{"src": "https://cdn-icons-png.flaticon.com/512/2907/2907253.png", "sizes": "512x512", "type": "image/png"}]
    })

@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request): return templates.TemplateResponse(request=request, name="login.html")

@app.post("/login")
def login_user(request: Request, response: Response, username: str=Form(...), password: str=Form(...), db: Session=Depends(get_db)):
    user = db.query(models.Employee).filter(models.Employee.username == username).first()
    if not user or user.password_hash != hash_password(password):
        return templates.TemplateResponse("login.html", {"request": request, "error": "Неверный логин или пароль"})
    redirect = RedirectResponse(url="/", status_code=303)
    redirect.set_cookie(key="user_id", value=str(user.id), httponly=True)
    return redirect

@app.get("/logout")
def logout():
    response = RedirectResponse(url="/login")
    response.delete_cookie("user_id")
    return response

@app.get("/", response_class=HTMLResponse)
def get_employees_page(request: Request, search: Optional[str] = None, sort: Optional[str] = "name", user_id: Optional[str] = Cookie(None), db: Session = Depends(get_db)):
    if not user_id: return RedirectResponse(url="/login")
    current_user = db.query(models.Employee).filter(models.Employee.id == int(user_id)).first()
    if not current_user:
        response = RedirectResponse(url="/login")
        response.delete_cookie("user_id")
        return response
    
    unread_notifs = db.query(models.Notification).filter(models.Notification.employee_id == current_user.id, models.Notification.is_read == False).all()
    user_schedule = current_user.schedule.split()
    working_today = user_schedule[TODAY_INDEX] if len(user_schedule) > TODAY_INDEX and user_schedule[TODAY_INDEX] in ['Д', 'Н'] else None

    if current_user.role in ["Админ", "Менеджер"]:
        query = db.query(models.Employee)
        if search: query = query.filter(models.Employee.name.icontains(search))
        if sort == "role": query = query.order_by(models.Employee.role.asc(), models.Employee.name.asc())
        else: query = query.order_by(models.Employee.name.asc())
        employees_data = query.all()
        pending_requests = db.query(models.LeaveRequest).filter(models.LeaveRequest.status == "Ожидает").all()
    else:
        employees_data = [current_user]
        pending_requests = []

    # --- НОВЫЙ ПОДСЧЕТ ПО ЧАСАМ ---
    for emp in employees_data:
        shifts = emp.schedule.split()
        work_hours = 0
        for s in shifts:
            if s in ['Д', 'Н']:
                work_hours += 12  # Дневная и ночная смены по 12 часов
            elif s == 'Р':
                work_hours += 8   # Обычная смена 8 часов
                
        emp.work_hours = work_hours
        emp.salary = work_hours * emp.hourly_rate

    return templates.TemplateResponse(
        request=request, name="employees.html", 
        context={"current_user": current_user, "employees": employees_data, "days": DAYS, "pending_requests": pending_requests, 
                 "search": search or "", "sort": sort, "unread_notifs": unread_notifs, "working_today": working_today, "today_name": DAYS[TODAY_INDEX]}
    )

@app.get("/read-notifications")
def read_notifications(user_id: Optional[str] = Cookie(None), db: Session = Depends(get_db)):
    if user_id:
        db.query(models.Notification).filter(models.Notification.employee_id == int(user_id)).update({"is_read": True})
        db.commit()
    return RedirectResponse(url="/", status_code=303)

@app.post("/request-leave/{emp_id}/{day_index}")
def request_leave_day(emp_id: int, day_index: int, user_id: Optional[str] = Cookie(None), db: Session = Depends(get_db)):
    if user_id:
        req = models.LeaveRequest(employee_id=emp_id, day_index=day_index)
        db.add(req)
        db.commit()
    return RedirectResponse(url="/", status_code=303)

@app.post("/process-leave/{req_id}/{action}")
def process_leave(req_id: int, action: str, user_id: Optional[str] = Cookie(None), db: Session = Depends(get_db)):
    current_user = db.query(models.Employee).filter(models.Employee.id == int(user_id)).first()
    if current_user and current_user.role in ["Админ", "Менеджер"]:
        req = db.query(models.LeaveRequest).filter(models.LeaveRequest.id == req_id).first()
        if req:
            emp = req.employee
            if action == "approve":
                req.status = "Одобрено"
                schedule_list = emp.schedule.split()
                if req.day_index < len(schedule_list): schedule_list[req.day_index] = "О"
                emp.schedule = " ".join(schedule_list)
                notif_msg = f"🟢 Ваш отгул на {DAYS[req.day_index]} ОДОБРЕН."
            else:
                req.status = "Отклонено"
                notif_msg = f"🔴 Ваш отгул на {DAYS[req.day_index]} ОТКЛОНЕН."
            db.add(models.Notification(employee_id=emp.id, message=notif_msg))
            db.commit()
    return RedirectResponse(url="/", status_code=303)

@app.post("/apply-template/{emp_id}/{template}")
def apply_template(emp_id: int, template: str, user_id: Optional[str] = Cookie(None), db: Session = Depends(get_db)):
    current_user = db.query(models.Employee).filter(models.Employee.id == int(user_id)).first()
    if current_user and current_user.role in ["Админ", "Менеджер"]:
        emp = db.query(models.Employee).filter(models.Employee.id == emp_id).first()
        if emp:
            if template == "2_2":
                pattern = ["Д", "Д", "В", "В"]
                emp.schedule = " ".join([pattern[i % 4] for i in range(30)])
            elif template == "5_2":
                pattern = ["Д", "Д", "Д", "Д", "Д", "В", "В"]
                emp.schedule = " ".join([pattern[i % 7] for i in range(30)])
            db.commit()
    return RedirectResponse(url="/", status_code=303)

@app.post("/update-shift/{emp_id}/{day_index}")
def update_shift(emp_id: int, day_index: int, new_shift: str = Form(...), user_id: Optional[str] = Cookie(None), db: Session = Depends(get_db)):
    current_user = db.query(models.Employee).filter(models.Employee.id == int(user_id)).first() if user_id else None
    if not current_user or current_user.role not in ["Админ", "Менеджер"]: return RedirectResponse(url="/", status_code=303)
    emp = db.query(models.Employee).filter(models.Employee.id == emp_id).first()
    if emp:
        schedule_list = emp.schedule.split()
        if day_index < len(schedule_list):
            schedule_list[day_index] = new_shift
            emp.schedule = " ".join(schedule_list)
        db.commit()
    return RedirectResponse(url="/", status_code=303)
    
@app.post("/add-employee")
def add_employee(name: str=Form(...), role: str=Form(...), db: Session=Depends(get_db)):
    last_emp = db.query(models.Employee).order_by(models.Employee.id.desc()).first()
    new_id = (last_emp.id + 1) if last_emp else 1
    new_emp = models.Employee(name=name, role=role, username=f"user{new_id}", password_hash=hash_password("1234"), schedule=" ".join(["В"]*30))
    db.add(new_emp)
    db.commit()
    return RedirectResponse(url="/", status_code=303)

@app.post("/update-profile")
def update_profile(phone: str=Form(""), email: str=Form(""), new_password: str=Form(""), user_id: Optional[str]=Cookie(None), db: Session=Depends(get_db)):
    user = db.query(models.Employee).filter(models.Employee.id == int(user_id)).first()
    if user:
        user.phone, user.email = phone, email
        if new_password.strip(): user.password_hash = hash_password(new_password)
        db.commit()
    return RedirectResponse(url="/", status_code=303)

@app.post("/reset-password/{emp_id}")
def reset_password(emp_id: int, user_id: Optional[str]=Cookie(None), db: Session=Depends(get_db)):
    current_user = db.query(models.Employee).filter(models.Employee.id == int(user_id)).first() if user_id else None
    if current_user and current_user.role == "Админ":
        emp = db.query(models.Employee).filter(models.Employee.id == emp_id).first()
        if emp:
            emp.password_hash = hash_password("1234")
            db.commit()
    return RedirectResponse(url="/", status_code=303)

@app.get("/export-calendar")
def export_calendar(user_id: Optional[str] = Cookie(None), db: Session = Depends(get_db)):
    user = db.query(models.Employee).filter(models.Employee.id == int(user_id)).first()
    ics_content = ["BEGIN:VCALENDAR", "VERSION:2.0", "PRODID:-//Company Schedule//RU"]
    for i, shift in enumerate(user.schedule.split()):
        if shift in ["Д", "Н", "Р"]:
            d_str = f"202611{i+1:02d}"
            ics_content.extend(["BEGIN:VEVENT", f"DTSTART:{d_str}T090000Z", f"DTEND:{d_str}T180000Z", f"SUMMARY:Смена ({shift})", "END:VEVENT"])
    ics_content.append("END:VCALENDAR")
    return Response(content="\r\n".join(ics_content), media_type="text/calendar", headers={"Content-Disposition": f"attachment; filename=schedule.ics"})

@app.get("/export-excel")
def export_excel(user_id: Optional[str] = Cookie(None), db: Session = Depends(get_db)):
    employees = db.query(models.Employee).order_by(models.Employee.name.asc()).all()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ID", "ФИО", "Должность", "Телефон", "Email"] + DAYS)
    for emp in employees:
        ws.append([emp.id, emp.name, emp.role, emp.phone or "", emp.email or ""] + emp.schedule.split()[:len(DAYS)])
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return StreamingResponse(output, headers={'Content-Disposition': 'attachment; filename="company_schedule.xlsx"'}, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.post("/import-excel")
def import_excel(file: UploadFile = File(...), user_id: Optional[str] = Cookie(None), db: Session = Depends(get_db)):
    contents = file.file.read()
    ws = openpyxl.load_workbook(filename=io.BytesIO(contents), data_only=True).active
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[1]: continue
        emp_id, name, role, phone, email = row[0], str(row[1]), str(row[2] or "Сотрудник"), str(row[3] or ""), str(row[4] or "")
        schedule_str = " ".join([str(cell).upper().strip() if cell else "В" for cell in row[5:35]])
        
        emp = db.query(models.Employee).filter(models.Employee.id == int(emp_id)).first() if emp_id else db.query(models.Employee).filter(models.Employee.name == name).first()
        if emp:
            emp.name, emp.role, emp.phone, emp.email, emp.schedule = name, role, phone, email, schedule_str
        else:
            last = db.query(models.Employee).order_by(models.Employee.id.desc()).first()
            db.add(models.Employee(name=name, role=role, phone=phone, email=email, schedule=schedule_str, username=f"user{(last.id + 1) if last else 1}", password_hash=hash_password("1234")))
    db.commit()
    return RedirectResponse(url="/", status_code=303)